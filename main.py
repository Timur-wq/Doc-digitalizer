import sys
import requests
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image
from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QVBoxLayout, QPushButton,
    QLabel, QFileDialog, QWidget, QProgressBar, QMessageBox,
    QHBoxLayout, QScrollArea, QCheckBox
)
from PyQt5.QtCore import Qt
from PyQt5.QtGui import QPixmap

import re
import time

class OCRApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.initUI()
        self.selected_file = None
        self.image_width = 0
        self.image_height = 0

    def initUI(self):
        self.setWindowTitle('OCR Приложение')
        self.setGeometry(100, 100, 800, 600)

        main_layout = QHBoxLayout()

        # Левая панель управления
        control_panel = QVBoxLayout()

        self.label = QLabel('Выберите изображение для распознавания текста')
        self.label.setAlignment(Qt.AlignCenter)
        control_panel.addWidget(self.label)

        self.select_btn = QPushButton('Выбрать файл')
        self.select_btn.clicked.connect(self.select_file)
        control_panel.addWidget(self.select_btn)

        self.handwritten_checkbox = QCheckBox('Рукописный текст')
        control_panel.addWidget(self.handwritten_checkbox)

        self.process_btn = QPushButton('Распознать и сохранить')
        self.process_btn.clicked.connect(self.process_image)
        self.process_btn.setEnabled(False)
        control_panel.addWidget(self.process_btn)

        self.progress = QProgressBar()
        self.progress.setVisible(False)
        control_panel.addWidget(self.progress)

        # Правая панель с изображением
        self.image_label = QLabel()
        self.image_label.setAlignment(Qt.AlignCenter)
        self.image_label.setStyleSheet("border: 1px solid black;")

        scroll = QScrollArea()
        scroll.setWidget(self.image_label)
        scroll.setWidgetResizable(True)

        # Компоновка
        control_widget = QWidget()
        control_widget.setLayout(control_panel)
        control_widget.setMaximumWidth(300)

        main_layout.addWidget(control_widget)
        main_layout.addWidget(scroll)

        container = QWidget()
        container.setLayout(main_layout)
        self.setCentralWidget(container)

    def select_file(self):
        options = QFileDialog.Options()
        file_name, _ = QFileDialog.getOpenFileName(
            self, "Выберите изображение", "",
            "Images (*.png *.jpg *.jpeg *.bmp);;All Files (*)", options=options
        )

        if file_name:
            self.selected_file = file_name
            self.label.setText(f"Выбран файл: {file_name.split('/')[-1]}")
            self.process_btn.setEnabled(True)

            # Загрузка изображения для предпросмотра
            pixmap = QPixmap(file_name)
            if not pixmap.isNull():
                pixmap = pixmap.scaled(1000, 800, Qt.KeepAspectRatio)
                self.image_label.setPixmap(pixmap)
            else:
                self.image_label.setText("Ошибка загрузки изображения")

            # Получение реальных размеров изображения
            try:
                with Image.open(file_name) as img:
                    self.image_width, self.image_height = img.size
            except Exception as e:
                QMessageBox.warning(self, 'Ошибка', f'Ошибка получения размеров: {str(e)}')
                self.selected_file = None
                self.process_btn.setEnabled(False)

    def process_image(self):
        if not self.selected_file:
            QMessageBox.warning(self, 'Ошибка', 'Файл не выбран!')
            return

        self.progress.setVisible(True)
        self.progress.setValue(0)

        try:
            ocr_start = time.time()
            # Выполнение OCR
            self.progress.setValue(30)
            ocr_result = self.ocr_space_file(self.selected_file)
            ocr_end = time.time()
            print(f"Время выполнения OCR: {ocr_end - ocr_start:.2f} секунд")

            # Сохранение документа
            self.progress.setValue(60)
            output_filename, _ = QFileDialog.getSaveFileName(
                self, "Сохранить как Word документ", "",
                "Word Documents (*.docx);;All Files (*)"
            )

            if output_filename:
                if not output_filename.endswith('.docx'):
                    output_filename += '.docx'
                doc_start = time.time()
                self.create_word_doc(
                    ocr_result,
                    output_filename,
                    self.image_width
                )
                doc_end = time.time()
                print(f"Время создания Word-документа: {doc_end - doc_start:.2f} секунд")
                self.progress.setValue(100)
                QMessageBox.information(
                    self, 'Успех',
                    f'Документ сохранён как {output_filename}'
                )

        except Exception as e:
            QMessageBox.critical(self, 'Ошибка', f'Ошибка обработки: {str(e)}')
        finally:
            self.progress.setVisible(False)

    def ocr_space_file(self, filename, overlay=True):
        payload = {
            'isOverlayRequired': overlay,
            'apikey': 'K86743096388957',
            'language': 'rus',
        }

        if self.handwritten_checkbox.isChecked():
            payload.update({'ocrengine': 5, 'isTable': True, 'detectOrientation': True, 'isOverlayRequired': overlay})

        with open(filename, 'rb') as f:
            response = requests.post(
                'https://api.ocr.space/parse/image',
                files={filename: f},
                data=payload,
            )
        return response.json()

    def create_word_doc(self, ocr_result, output_filename, image_width):
        doc = Document()

        try:
            parsed_data = ocr_result['ParsedResults'][0]
            lines_data = parsed_data['TextOverlay']['Lines']
        except (KeyError, IndexError):
            QMessageBox.warning(self, 'Ошибка', 'Некорректные данные OCR')
            return

        # Собираем все текстовые строки
        all_lines = []
        for line in lines_data:
            line_text = line.get('LineText', '').strip()
            if line_text:
                all_lines.append({
                    'text': line_text,
                    'words': line.get('Words', []),
                    'line_obj': line
                })

        # Расширенный шаблон для служебных строк
        service_pattern = r'''
            \(
            (?:серия\s*[,.!]?\s*номер\s*[,.!]?\s*кем\s*и\s*когда\s*выдан|
            (?:ф[.!]?и[.!]?о\.?)(?:\s*полностью)?|
            адрес\s*(?:регистрации|проживания)|
            дата\s*рождения|
            подпись)
            \)
        '''.replace('\n', '').replace(' ', r'\s*')

        # Обработка строк
        current_block = []  # Для группировки рукописного текста
        i = 0
        while i < len(all_lines):
            line_data = all_lines[i]
            current_text = line_data['text']
            cleaned_text = re.sub(r'_+', '', current_text)
            normalized_text = re.sub(r'[!]', '.', cleaned_text.lower())

            # Пропуск служебных строк
            if re.search(service_pattern, normalized_text, re.IGNORECASE | re.VERBOSE):
                i += 1
                continue

            # Обнаружение начала рукописного блока
            if '_' in current_text:
                # Собираем все последующие строки до следующего служебного шаблона
                current_block = []
                i += 1  # Пропускаем строку с подчеркиванием
                while i < len(all_lines):
                    next_line = all_lines[i]
                    next_text = re.sub(r'_+', '', next_line['text'])
                    next_normalized = re.sub(r'[!]', '.', next_text.lower())

                    # Проверка на конец блока
                    if re.search(service_pattern, next_normalized, re.IGNORECASE | re.VERBOSE):
                        break

                    if next_text.strip():
                        current_block.append(next_text.strip())
                    i += 1

                # Добавляем собранный блок как единый абзац
                if current_block:
                    p = doc.add_paragraph(' '.join(current_block))
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT  # Выравнивание по левому краю
                continue

            # Обработка обычного текста
            if cleaned_text.strip():
                p = doc.add_paragraph(cleaned_text)
                # Выравнивание по координатам
                try:
                    words = line_data['words']
                    left = min(word['Left'] for word in words)
                    right = max(word['Left'] + word['Width'] for word in words)
                    center_x = (left + right) / 2
                    if abs(center_x - image_width / 2) <= image_width * 0.1:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                except:
                    pass

            i += 1

        doc.save(output_filename)


if __name__ == '__main__':
    app = QApplication(sys.argv)
    ex = OCRApp()
    ex.show()
    sys.exit(app.exec_())